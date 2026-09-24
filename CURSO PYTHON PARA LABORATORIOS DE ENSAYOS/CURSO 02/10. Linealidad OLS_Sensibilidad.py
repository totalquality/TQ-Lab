# ======================================================================
# LINEALIDAD + PLAN B POLINOMIAL + SENSIBILIDAD
# Validación de métodos químicos cuantitativos
#
# Criterios y supuestos OLS mantenidos tal como el script inicial:
#   1. Significancia de la pendiente / modelo
#   2. Lack-of-Fit + Error puro
#   3. R², R² ajustado y R² predicho
#   4. Normalidad de residuales: Anderson-Darling
#   5. Independencia: Residuales vs orden + Durbin-Watson
#   6. Homocedasticidad: Residuales vs ajustados + Breusch-Pagan
#   7. Diagnóstico de observaciones: residual estandarizado,
#      leverage y distancia de Cook
#
# PLAN B:
#   Si el modelo lineal presenta Lack-of-Fit significativo,
#   se evalúa polinomio de grado 2 y luego de grado 3.
#
# SENSIBILIDAD:
#   OLS: pendiente.
#   Polinomio: dY/dX.
# ======================================================================

import warnings
warnings.filterwarnings("ignore")

from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import statsmodels.api as sm

from scipy import stats
from statsmodels.stats.diagnostic import normal_ad, het_breuschpagan
from statsmodels.stats.stattools import durbin_watson


# ======================================================================
# 1. CONFIGURACIÓN
# ======================================================================

ARCHIVO = Path(
    r"D:\Ciencia_de_Datos\Validación de métodos químicos\linealidad_sensibilidad.xlsx"
)

HOJA = None
ALPHA = 0.05

ARCHIVO_SALIDA = Path(
    r"D:\Ciencia_de_Datos\Validación de métodos químicos"
    r"\resultado_linealidad_sensibilidad_CORREGIDO.xlsx"
)


# ======================================================================
# 2. ESTILO VISUAL — MISMA LÓGICA DEL SCRIPT INICIAL
# ======================================================================

plt.rcParams.update({
    "figure.dpi": 120,
    "savefig.dpi": 220,
    "font.family": "DejaVu Sans",
    "axes.titlesize": 14,
    "axes.titleweight": "bold",
    "axes.labelsize": 11,
    "axes.edgecolor": "#263238",
    "axes.linewidth": 1.0,
    "xtick.labelsize": 9.5,
    "ytick.labelsize": 9.5,
    "legend.fontsize": 9.5,
    "figure.facecolor": "white",
    "axes.facecolor": "#FBFCFE",
})


# ======================================================================
# 3. FUNCIONES AUXILIARES
# ======================================================================

def fmt(x, d=6):
    if x is None:
        return "NA"
    try:
        if not np.isfinite(x):
            return "NA"
    except Exception:
        return "NA"
    return f"{x:.{d}g}"


def fmt_p(p):
    if p is None:
        return "NA"
    try:
        if not np.isfinite(p):
            return "NA"
    except Exception:
        return "NA"
    if p < 0.000001:
        return "< 0.000001"
    return f"{p:.6f}"


def resultado_p(p, criterio):
    if not np.isfinite(p):
        return "NO EVALUABLE"
    if criterio == "menor":
        return "CUMPLE" if p < ALPHA else "REVISAR"
    return "CUMPLE" if p >= ALPHA else "REVISAR"


def construir_X(x, grado=1):
    """Matriz de diseño para polinomio de grado 1, 2 o 3."""
    if grado == 1:
        return sm.add_constant(np.asarray(x, dtype=float))
    return sm.add_constant(
        np.column_stack([
            np.asarray(x, dtype=float) ** i
            for i in range(1, grado + 1)
        ])
    )


def predecir_modelo(modelo, x, grado):
    return modelo.predict(construir_X(x, grado))


def calcular_r2_predicho(modelo):
    """R² predicho mediante PRESS, como en el script inicial."""
    influencia = modelo.get_influence()
    resid_press = influencia.resid_press
    press = np.sum(resid_press ** 2)
    tss = np.sum((modelo.model.endog - np.mean(modelo.model.endog)) ** 2)
    r2_pred = 1 - press / tss if tss > 0 else np.nan
    return press, r2_pred


def lack_of_fit(x, y, y_fitted, n_parameters):
    """
    SSE = SS_LOF + SS_PE

    SS_PE  = error puro dentro de los niveles replicados.
    SS_LOF = SSE - SS_PE.

    p = número de parámetros del modelo:
      lineal     -> 2
      cuadrático -> 3
      cúbico     -> 4
    """
    tabla = pd.DataFrame({
        "X": x,
        "Y": y,
        "Y_fit": y_fitted
    })

    grupos = tabla.groupby("X", sort=True)
    niveles = grupos.ngroups
    n_total = len(tabla)

    sspe = 0.0
    for _, grupo in grupos:
        media = grupo["Y"].mean()
        sspe += np.sum((grupo["Y"] - media) ** 2)

    sse = np.sum((y - y_fitted) ** 2)
    sslof = sse - sspe

    if sslof < 0 and abs(sslof) < 1e-10:
        sslof = 0.0

    df_lof = niveles - n_parameters
    df_pe = n_total - niveles

    if df_lof <= 0 or df_pe <= 0 or sspe <= 0:
        return {
            "evaluable": False,
            "niveles": niveles,
            "N": n_total,
            "SSPE": sspe,
            "SSLOF": sslof,
            "MSLOF": np.nan,
            "MSPE": np.nan,
            "F_LOF": np.nan,
            "p_LOF": np.nan,
            "df_LOF": df_lof,
            "df_PE": df_pe,
        }

    mslof = sslof / df_lof
    mspe = sspe / df_pe
    f_lof = mslof / mspe
    p_lof = stats.f.sf(f_lof, df_lof, df_pe)

    return {
        "evaluable": True,
        "niveles": niveles,
        "N": n_total,
        "SSPE": sspe,
        "SSLOF": sslof,
        "MSLOF": mslof,
        "MSPE": mspe,
        "F_LOF": f_lof,
        "p_LOF": p_lof,
        "df_LOF": df_lof,
        "df_PE": df_pe,
    }


def evaluar_supuestos(modelo):
    """Exactamente las tres pruebas del script inicial."""
    residuos = np.asarray(modelo.resid, dtype=float)
    ajustados = np.asarray(modelo.fittedvalues, dtype=float)

    try:
        ad, p_ad = normal_ad(residuos)
    except Exception:
        ad, p_ad = np.nan, np.nan

    try:
        bp_lm, bp_p_lm, bp_f, bp_p_f = het_breuschpagan(
            residuos,
            modelo.model.exog
        )
    except Exception:
        bp_lm, bp_p_lm, bp_f, bp_p_f = (
            np.nan, np.nan, np.nan, np.nan
        )

    try:
        dw = durbin_watson(residuos)
    except Exception:
        dw = np.nan

    return {
        "AD": ad,
        "p_AD": p_ad,
        "BP_LM": bp_lm,
        "BP_p_LM": bp_p_lm,
        "BP_F": bp_f,
        "BP_p_F": bp_p_f,
        "DW": dw,
        "residuos": residuos,
        "ajustados": ajustados,
    }


def obtener_diagnosticos(modelo, x, y, grado):
    """
    Mismos diagnósticos del script inicial:
    residual estandarizado, leverage y Cook.
    """
    influencia = modelo.get_influence()

    datos_d = pd.DataFrame({
        "Obs": np.arange(1, len(x) + 1),
        "X": x,
        "Y": y,
        "Ajuste": modelo.fittedvalues,
        "Resid": modelo.resid,
        "Resid est": influencia.resid_studentized_internal,
        "Resid est externo": influencia.resid_studentized_external,
        "Leverage": influencia.hat_matrix_diag,
        "Cook": influencia.cooks_distance[0],
        "DFFITS": influencia.dffits[0],
    })

    try:
        pred = modelo.get_prediction(construir_X(x, grado))
        frame = pred.summary_frame(alpha=ALPHA)

        datos_d["EE ajuste"] = frame["mean_se"].values
        datos_d["IC95 media inferior"] = frame["mean_ci_lower"].values
        datos_d["IC95 media superior"] = frame["mean_ci_upper"].values
        datos_d["IP95 inferior"] = frame["obs_ci_lower"].values
        datos_d["IP95 superior"] = frame["obs_ci_upper"].values
    except Exception:
        datos_d["EE ajuste"] = np.nan
        datos_d["IC95 media inferior"] = np.nan
        datos_d["IC95 media superior"] = np.nan
        datos_d["IP95 inferior"] = np.nan
        datos_d["IP95 superior"] = np.nan

    n = len(x)
    p = grado + 1

    umbral_resid = 2
    umbral_leverage = 2 * p / n
    umbral_cook = 4 / n

    datos_d["Marca R"] = np.where(
        np.abs(datos_d["Resid est"]) > umbral_resid, "R", ""
    )
    datos_d["Marca leverage"] = np.where(
        datos_d["Leverage"] > umbral_leverage, "L", ""
    )
    datos_d["Marca Cook"] = np.where(
        datos_d["Cook"] > umbral_cook, "C", ""
    )

    return datos_d, umbral_resid, umbral_leverage, umbral_cook


def ajustar_modelo(x, y, grado):
    modelo = sm.OLS(y, construir_X(x, grado)).fit()
    press, r2_pred = calcular_r2_predicho(modelo)
    lof = lack_of_fit(
        x, y, modelo.fittedvalues, n_parameters=grado + 1
    )
    sup = evaluar_supuestos(modelo)
    diag = obtener_diagnosticos(modelo, x, y, grado)

    return {
        "grado": grado,
        "modelo": modelo,
        "press": press,
        "r2_predicho": r2_pred,
        "lof": lof,
        "supuestos": sup,
        "diagnosticos": diag,
    }


def ecuacion_modelo(modelo, grado):
    b = np.asarray(modelo.params)

    if grado == 1:
        signo = "+" if b[1] >= 0 else "-"
        return (
            f"Y = {b[0]:.10g} {signo} "
            f"{abs(b[1]):.10g}X"
        )

    if grado == 2:
        s2 = "+" if b[2] >= 0 else "-"
        return (
            f"Y = {b[0]:.10g} "
            f"{'+' if b[1] >= 0 else '-'} {abs(b[1]):.10g}X "
            f"{s2} {abs(b[2]):.10g}X²"
        )

    s2 = "+" if b[2] >= 0 else "-"
    s3 = "+" if b[3] >= 0 else "-"
    return (
        f"Y = {b[0]:.10g} "
        f"{'+' if b[1] >= 0 else '-'} {abs(b[1]):.10g}X "
        f"{s2} {abs(b[2]):.10g}X² "
        f"{s3} {abs(b[3]):.10g}X³"
    )


def sensibilidad_grid(modelo, grado, x_grid):
    b = np.asarray(modelo.params)

    if grado == 1:
        return np.full_like(x_grid, b[1], dtype=float)

    if grado == 2:
        return b[1] + 2 * b[2] * x_grid

    return b[1] + 2 * b[2] * x_grid + 3 * b[3] * x_grid**2


def evaluar_polynomial_candidate(resultado):
    """
    Criterios funcionales del Plan B:
      - modelo significativo
      - término de mayor grado significativo
      - no evidencia de Lack-of-Fit

    Los supuestos se reportan por separado con las mismas pruebas
    del script inicial y NO se confunden con falta de linealidad.
    """
    modelo = resultado["modelo"]
    grado = resultado["grado"]
    lof = resultado["lof"]

    p_termino = modelo.pvalues[grado]

    return (
        modelo.f_pvalue < ALPHA
        and p_termino < ALPHA
        and (
            not lof["evaluable"]
            or lof["p_LOF"] >= ALPHA
        )
    )


# ======================================================================
# 4. CARGAR DATOS
# ======================================================================

if not ARCHIVO.exists():
    raise FileNotFoundError(
        f"\nNo se encontró el archivo:\n{ARCHIVO}\n"
        "\nVerifique la ruta y el nombre."
    )

xls = pd.ExcelFile(ARCHIVO)
hoja_real = xls.sheet_names[0] if HOJA is None else HOJA

if hoja_real not in xls.sheet_names:
    raise ValueError(
        f"La hoja '{hoja_real}' no existe. "
        f"Hojas disponibles: {xls.sheet_names}"
    )

raw = pd.read_excel(ARCHIVO, sheet_name=hoja_real)

numericas = []
for col in raw.columns:
    serie = pd.to_numeric(raw[col], errors="coerce")
    if serie.notna().sum() >= 3:
        numericas.append(col)

if len(numericas) < 2:
    raise ValueError(
        "No se encontraron al menos dos columnas numéricas."
    )

col_x = numericas[0]
col_y = numericas[1]

base = pd.DataFrame({
    "X": pd.to_numeric(raw[col_x], errors="coerce"),
    "Y": pd.to_numeric(raw[col_y], errors="coerce")
}).dropna().reset_index(drop=True)

x = base["X"].to_numpy(float)
y = base["Y"].to_numpy(float)

if len(x) < 4:
    raise ValueError("Se necesitan más observaciones para la regresión.")

niveles_x = np.unique(x)

print("\n" + "=" * 95)
print("VALIDACIÓN — LINEALIDAD + PLAN B POLINOMIAL + SENSIBILIDAD")
print("=" * 95)
print(f"\nArchivo : {ARCHIVO}")
print(f"Hoja    : {hoja_real}")
print(f"X       : {col_x}")
print(f"Y       : {col_y}")
print(f"N       : {len(x)}")
print(f"Niveles : {len(niveles_x)}")
print(f"Valores X: {niveles_x}")


# ======================================================================
# 5. MODELO OLS LINEAL
# ======================================================================

res_ols = ajustar_modelo(x, y, grado=1)
modelo_ols = res_ols["modelo"]
lof_ols = res_ols["lof"]
sup_ols = res_ols["supuestos"]
diag_ols = res_ols["diagnosticos"]

intercepto = modelo_ols.params[0]
pendiente = modelo_ols.params[1]
p_intercepto = modelo_ols.pvalues[0]
p_pendiente = modelo_ols.pvalues[1]

R2 = modelo_ols.rsquared
R2_ajustado = modelo_ols.rsquared_adj
S = np.sqrt(modelo_ols.mse_resid)
PRESS = res_ols["press"]
R2_predicho = res_ols["r2_predicho"]


# ======================================================================
# 6. ECUACIÓN Y COEFICIENTES
# ======================================================================

print("\n" + "=" * 95)
print("1. ECUACIÓN DE REGRESIÓN")
print("=" * 95)
print(f"\n{ecuacion_modelo(modelo_ols, 1)}")

tabla_coef_ols = pd.DataFrame({
    "Término": ["Constante", str(col_x)],
    "Coef": [modelo_ols.params[0], modelo_ols.params[1]],
    "EE del coef.": [modelo_ols.bse[0], modelo_ols.bse[1]],
    "Valor T": [modelo_ols.tvalues[0], modelo_ols.tvalues[1]],
    "Valor p": [modelo_ols.pvalues[0], modelo_ols.pvalues[1]],
    "FIV": [np.nan, 1.0],
    "IC95 inferior": [
        modelo_ols.conf_int(alpha=ALPHA)[0, 0],
        modelo_ols.conf_int(alpha=ALPHA)[1, 0]
    ],
    "IC95 superior": [
        modelo_ols.conf_int(alpha=ALPHA)[0, 1],
        modelo_ols.conf_int(alpha=ALPHA)[1, 1]
    ]
})

print("\n" + "-" * 95)
print("2. COEFICIENTES DEL MODELO")
print("-" * 95)
print(tabla_coef_ols.to_string(index=False, float_format=lambda z: f"{z:.6g}"))


# ======================================================================
# 7. RESUMEN DEL MODELO
# ======================================================================

print("\n" + "-" * 95)
print("3. RESUMEN DEL MODELO")
print("-" * 95)
print(f"S                    = {S:.10g}")
print(f"R-cuadrado           = {R2 * 100:.8f} %")
print(f"R-cuadrado ajustado  = {R2_ajustado * 100:.8f} %")
print(f"R-cuadrado predicho  = {R2_predicho * 100:.8f} %")
print(f"PRESS                = {PRESS:.10g}")
print(
    "\nNota: no existe un umbral universal de R²; "
    "debe interpretarse junto con Lack-of-Fit y residuales."
)


# ======================================================================
# 8. ANOVA DE REGRESIÓN
# ======================================================================

df_reg = 1
df_error = len(x) - 2
df_total = len(x) - 1

ss_reg = modelo_ols.ess
ss_error = modelo_ols.ssr
ss_total = modelo_ols.centered_tss
ms_reg = ss_reg
ms_error = ss_error / df_error
f_reg = ms_reg / ms_error
p_reg = stats.f.sf(f_reg, df_reg, df_error)

tabla_anova_ols = pd.DataFrame({
    "Fuente": ["Regresión", "X", "Error", "Total"],
    "GL": [df_reg, 1, df_error, df_total],
    "SC": [ss_reg, ss_reg, ss_error, ss_total],
    "MC": [ms_reg, ms_reg, ms_error, np.nan],
    "F": [f_reg, f_reg, np.nan, np.nan],
    "p": [p_reg, p_reg, np.nan, np.nan]
})

print("\n" + "=" * 95)
print("4. ANÁLISIS DE VARIANZA DE LA REGRESIÓN")
print("=" * 95)
print(tabla_anova_ols.to_string(index=False, float_format=lambda z: f"{z:.6g}"))

print("\nH₀: β1 = 0")
print("H₁: β1 ≠ 0")
print(f"F = {f_reg:.10g}")
print(f"p = {fmt_p(p_reg)}")
print(
    "Resultado: " +
    ("CUMPLE — relación lineal significativa."
     if p_reg < ALPHA
     else "REVISAR — no se evidencia relación lineal significativa.")
)


# ======================================================================
# 9. LACK-OF-FIT + ERROR PURO
# ======================================================================

print("\n" + "=" * 95)
print("4.1 LACK-OF-FIT — PRUEBA DE FALTA DE AJUSTE")
print("=" * 95)

tabla_lof_ols = None

if lof_ols["evaluable"]:
    tabla_lof_ols = pd.DataFrame({
        "Fuente": ["Error", "  Falta de ajuste", "  Error puro"],
        "GL": [
            len(x) - 2,
            lof_ols["df_LOF"],
            lof_ols["df_PE"]
        ],
        "SC": [
            lof_ols["SSLOF"] + lof_ols["SSPE"],
            lof_ols["SSLOF"],
            lof_ols["SSPE"]
        ],
        "MC": [
            (lof_ols["SSLOF"] + lof_ols["SSPE"]) / (len(x) - 2),
            lof_ols["MSLOF"],
            lof_ols["MSPE"]
        ],
        "F": [np.nan, lof_ols["F_LOF"], np.nan],
        "p": [np.nan, lof_ols["p_LOF"], np.nan]
    })

    print(tabla_lof_ols.to_string(index=False, float_format=lambda z: f"{z:.6g}"))
    print(f"\nF = {lof_ols['F_LOF']:.10g}")
    print(f"p = {fmt_p(lof_ols['p_LOF'])}")
    print("\nH₀: no existe falta de ajuste significativa.")
    print("H₁: existe falta de ajuste significativa.")

    if lof_ols["p_LOF"] >= ALPHA:
        print(
            "Conclusión: NO SE RECHAZA H₀. "
            "No se evidencia falta de ajuste significativa."
        )
    else:
        print(
            "Conclusión: SE RECHAZA H₀. "
            "Existe evidencia de falta de ajuste."
        )
else:
    print(
        "Lack-of-Fit NO EVALUABLE. "
        "Se requieren réplicas para estimar el error puro."
    )


# ======================================================================
# 10. DIAGNÓSTICO DE OBSERVACIONES POCO COMUNES
# ======================================================================

print("\n" + "-" * 95)
print("5. AJUSTES Y DIAGNÓSTICO DE OBSERVACIONES POCO COMUNES")
print("-" * 95)

# diagnosticos contiene una tupla: (tabla_diagnostico, umbral_resid,
# umbral_leverage, umbral_cook).
# Se extrae la tabla antes de seleccionar las columnas.
tabla_diag_ols, umbral_resid, umbral_leverage, umbral_cook = diag_ols

tabla_obs = tabla_diag_ols[[
    "Obs", "X", "Y", "Ajuste", "EE ajuste", "Resid",
    "Resid est", "Leverage", "Cook",
    "Marca R", "Marca leverage", "Marca Cook"
]]

print(tabla_obs.to_string(index=False, float_format=lambda z: f"{z:.6g}"))

print("\nCriterios diagnósticos utilizados:")
print(f"• R: |Residual estandarizado| > {umbral_resid:.1f}")
print(f"• L: Leverage > 2p/n = {umbral_leverage:.6g}")
print(f"• C: Cook > 4/n = {umbral_cook:.6g}")
print(
    "Estos son umbrales de señalización; "
    "no implican eliminar automáticamente una observación."
)


# ======================================================================
# 11. SUPUESTO 1 — NORMALIDAD
# ======================================================================

print("\n" + "-" * 95)
print("6. SUPUESTO OLS #1 — NORMALIDAD DE LOS RESIDUALES")
print("-" * 95)
print("Prueba: Anderson-Darling")
print(f"A² = {sup_ols['AD']:.6g}")
print(f"p  = {fmt_p(sup_ols['p_AD'])}")

if sup_ols["p_AD"] >= ALPHA:
    print("Conclusión: NO SE RECHAZA H₀ — normalidad compatible con los datos.")
else:
    print("Conclusión: SE RECHAZA H₀ — revisar normalidad de los residuales.")

print("Complementar con Q-Q Plot e histograma de residuales.")


# ======================================================================
# 12. SUPUESTO 2 — INDEPENDENCIA
# ======================================================================

print("\n" + "-" * 95)
print("7. SUPUESTO OLS #2 — INDEPENDENCIA DE LOS RESIDUALES")
print("-" * 95)
print(f"Durbin-Watson = {sup_ols['DW']:.6g}")
print(
    "Referencia utilizada por el script inicial: "
    "DW entre 1.5 y 2.5 como intervalo orientativo."
)

if 1.5 <= sup_ols["DW"] <= 2.5:
    print("Resultado: CUMPLE — sin señal importante de autocorrelación.")
else:
    print("Resultado: REVISAR — DW fuera del intervalo orientativo 1.5–2.5.")

print("Complementar con Residuales vs Orden.")


# ======================================================================
# 13. SUPUESTO 3 — HOMOCEDASTICIDAD
# ======================================================================

print("\n" + "-" * 95)
print("8. SUPUESTO OLS #3 — HOMOCEDASTICIDAD")
print("-" * 95)
print("Prueba: Breusch-Pagan")
print(f"LM = {sup_ols['BP_LM']:.6g}")
print(f"p (LM) = {fmt_p(sup_ols['BP_p_LM'])}")
print(f"F = {sup_ols['BP_F']:.6g}")
print(f"p (F) = {fmt_p(sup_ols['BP_p_F'])}")

if sup_ols["BP_p_F"] >= ALPHA:
    print("Conclusión: no se evidencia heterocedasticidad significativa.")
else:
    print("Conclusión: se evidencia heterocedasticidad significativa.")

print("Complementar con Residuales vs Ajustados.")


# ======================================================================
# 14. MATRIZ OLS — CRITERIOS EXACTOS DEL SCRIPT INICIAL
# ======================================================================

tabla_criterios_ols = pd.DataFrame({
    "Criterio": [
        "Pendiente significativa",
        "Falta de ajuste",
        "Normalidad AD",
        "Homocedasticidad BP",
        "Independencia DW"
    ],
    "Resultado": [
        p_pendiente,
        lof_ols["p_LOF"],
        sup_ols["p_AD"],
        sup_ols["BP_p_F"],
        sup_ols["DW"]
    ],
    "Referencia": [
        "p < 0.05",
        "p ≥ 0.05",
        "p ≥ 0.05",
        "p ≥ 0.05",
        "1.5 ≤ DW ≤ 2.5"
    ],
    "Evaluación": [
        "CUMPLE" if p_pendiente < ALPHA else "REVISAR",
        (
            "CUMPLE"
            if np.isfinite(lof_ols["p_LOF"]) and lof_ols["p_LOF"] >= ALPHA
            else "REVISAR"
            if np.isfinite(lof_ols["p_LOF"])
            else "NO EVALUABLE"
        ),
        "CUMPLE" if sup_ols["p_AD"] >= ALPHA else "REVISAR",
        "CUMPLE" if sup_ols["BP_p_F"] >= ALPHA else "REVISAR",
        "CUMPLE" if 1.5 <= sup_ols["DW"] <= 2.5 else "REVISAR"
    ]
})

print("\n" + "=" * 95)
print("9. MATRIZ FINAL — MODELO OLS LINEAL")
print("=" * 95)
print(tabla_criterios_ols.to_string(index=False, float_format=lambda z: f"{z:.6g}"))


# ======================================================================
# 15. DECISIÓN AUTOMÁTICA DEL PLAN B
# ======================================================================
#
# TAL COMO SOLICITASTE:
# - El Plan B se activa por falta de ajuste significativa.
# - Los incumplimientos de supuestos no se convierten automáticamente
#   en "no linealidad".
# ======================================================================

activar_plan_b = (
    lof_ols["evaluable"]
    and lof_ols["p_LOF"] < ALPHA
)

modelo_final = modelo_ols
grado_final = 1
nombre_modelo_final = "OLS lineal"
resultado_polinomio = None


# ======================================================================
# 16. PLAN B — POLINOMIO GRADO 2 Y GRADO 3
# ======================================================================

if activar_plan_b:

    print("\n" + "=" * 95)
    print("10. PLAN B — REGRESIÓN POLINOMIAL")
    print("=" * 95)

    print(
        "\nLa prueba de Lack-of-Fit del modelo lineal es significativa."
    )
    print(
        "Se evaluará primero grado 2 y, si no cumple, grado 3."
    )

    for grado in [2, 3]:

        resultado = ajustar_modelo(x, y, grado)
        modelo = resultado["modelo"]
        lof = resultado["lof"]
        sup = resultado["supuestos"]

        p_termino = modelo.pvalues[grado]

        print("\n" + "-" * 95)
        print(f"POLINOMIO DE GRADO {grado}")
        print("-" * 95)
        print(f"\nEcuación:\n{ecuacion_modelo(modelo, grado)}")

        print("\nResumen del modelo:")
        print(f"S = {np.sqrt(modelo.mse_resid):.10g}")
        print(f"R² = {modelo.rsquared * 100:.8f} %")
        print(f"R² ajustado = {modelo.rsquared_adj * 100:.8f} %")
        print(f"R² predicho = {resultado['r2_predicho'] * 100:.8f} %")
        print(f"p del modelo = {fmt_p(modelo.f_pvalue)}")

        print("\nCriterios funcionales:")
        print(
            f"• Significancia del modelo: p = {fmt_p(modelo.f_pvalue)} "
            f"→ {'CUMPLE' if modelo.f_pvalue < ALPHA else 'REVISAR'}"
        )
        print(
            f"• Término X^{grado}: p = {fmt_p(p_termino)} "
            f"→ {'CUMPLE' if p_termino < ALPHA else 'REVISAR'}"
        )

        if lof["evaluable"]:
            print(
                f"• Lack-of-Fit: p = {fmt_p(lof['p_LOF'])} "
                f"→ {'CUMPLE' if lof['p_LOF'] >= ALPHA else 'REVISAR'}"
            )
        else:
            print("• Lack-of-Fit: NO EVALUABLE")

        print("\nSupuestos — MISMAS PRUEBAS DEL SCRIPT INICIAL:")
        print(
            f"• Anderson-Darling: p = {fmt_p(sup['p_AD'])} "
            f"→ {'CUMPLE' if sup['p_AD'] >= ALPHA else 'REVISAR'}"
        )
        print(
            f"• Breusch-Pagan: p = {fmt_p(sup['BP_p_F'])} "
            f"→ {'CUMPLE' if sup['BP_p_F'] >= ALPHA else 'REVISAR'}"
        )
        print(
            f"• Durbin-Watson: {sup['DW']:.6g} "
            f"→ {'CUMPLE' if 1.5 <= sup['DW'] <= 2.5 else 'REVISAR'}"
        )

        adecuado = (
            modelo.f_pvalue < ALPHA
            and p_termino < ALPHA
            and (
                not lof["evaluable"]
                or lof["p_LOF"] >= ALPHA
            )
        )

        if adecuado:
            resultado_polinomio = resultado
            modelo_final = modelo
            grado_final = grado
            nombre_modelo_final = f"Regresión polinomial grado {grado}"

            print(
                f"\n✓ Se selecciona el polinomio de grado {grado} "
                "como primer modelo que cumple los criterios funcionales."
            )
            break

        print(
            f"\n✗ El polinomio de grado {grado} no cumple todos "
            "los criterios funcionales. Se continúa."
        )

    if resultado_polinomio is None:
        print(
            "\nNINGUNO de los polinomios evaluados cumplió todos "
            "los criterios funcionales."
        )
        print(
            "No se fuerza la selección de un modelo polinomial."
        )


# ======================================================================
# 17. SENSIBILIDAD — PÁGINAS 220-223
# ======================================================================

print("\n" + "=" * 95)
print("11. SENSIBILIDAD ANALÍTICA")
print("=" * 95)

print("\nSensibilidad = dY/dX")

if grado_final == 1:

    sensibilidad = modelo_final.params[1]

    print(f"\nModelo final: OLS lineal")
    print(f"Sensibilidad = pendiente = {sensibilidad:.10g}")

else:

    b = np.asarray(modelo_final.params)

    if grado_final == 2:
        formula = (
            f"S(X) = {b[1]:.10g} + "
            f"2({b[2]:.10g})X"
        )
    else:
        formula = (
            f"S(X) = {b[1]:.10g} + "
            f"2({b[2]:.10g})X + "
            f"3({b[3]:.10g})X²"
        )

    sensibilidad = sensibilidad_grid(
        modelo_final,
        grado_final,
        np.array([np.mean(x)])
    )[0]

    print(f"\nModelo final: {nombre_modelo_final}")
    print("La sensibilidad no es constante.")
    print(formula)
    print(f"Sensibilidad en X medio ({np.mean(x):.6g}) = {sensibilidad:.10g}")


# ======================================================================
# 18. GRÁFICA 1 — CURVA DE CALIBRACIÓN, COMO SCRIPT INICIAL
# ======================================================================

x_grid = np.linspace(np.min(x), np.max(x), 500)
pred_grid = modelo_final.get_prediction(
    construir_X(x_grid, grado_final)
)
frame_grid = pred_grid.summary_frame(alpha=ALPHA)

fig = plt.figure(figsize=(14, 8), constrained_layout=True)
gs = fig.add_gridspec(1, 2, width_ratios=[3.8, 1.35])

ax = fig.add_subplot(gs[0, 0])
info = fig.add_subplot(gs[0, 1])

ax.fill_between(
    x_grid,
    frame_grid["obs_ci_lower"],
    frame_grid["obs_ci_upper"],
    alpha=0.14,
    label="Intervalo de predicción 95 %"
)

ax.fill_between(
    x_grid,
    frame_grid["mean_ci_lower"],
    frame_grid["mean_ci_upper"],
    alpha=0.20,
    label="IC 95 % de la media"
)

ax.plot(
    x_grid,
    frame_grid["mean"],
    color="#8E44AD",
    linewidth=2.8,
    label=nombre_modelo_final,
    zorder=4
)

ax.scatter(
    x,
    y,
    s=54,
    color="#1F77B4",
    alpha=0.90,
    edgecolor="white",
    linewidth=0.8,
    label="Datos experimentales",
    zorder=5
)

medias_nivel = (
    base.groupby("X")["Y"]
    .agg(["mean", "count"])
    .reset_index()
)

ax.scatter(
    medias_nivel["X"],
    medias_nivel["mean"],
    s=115,
    marker="D",
    color="#E67E22",
    edgecolor="white",
    linewidth=1.0,
    label="Media por nivel",
    zorder=6
)

ax.set_title(
    "Validación de linealidad — modelo seleccionado",
    fontsize=17,
    fontweight="bold"
)
ax.set_xlabel(str(col_x), fontsize=11.5)
ax.set_ylabel(str(col_y), fontsize=11.5)
ax.grid(True, alpha=0.18, linewidth=0.8)
ax.spines["top"].set_visible(False)
ax.spines["right"].set_visible(False)
ax.legend(loc="best", frameon=True, fancybox=True)

info.axis("off")
info.text(
    0.02, 0.96,
    "RESULTADO ESTADÍSTICO",
    fontsize=14, fontweight="bold", va="top"
)

texto_info = (
    f"N = {len(x)}\n\n"
    f"Modelo = {nombre_modelo_final}\n\n"
    f"S = {np.sqrt(modelo_final.mse_resid):.6g}\n"
    f"R² = {modelo_final.rsquared * 100:.6f} %\n"
    f"R² ajustado = {modelo_final.rsquared_adj * 100:.6f} %\n"
    f"R² predicho = {(
        res_ols['r2_predicho'] if grado_final == 1
        else resultado_polinomio['r2_predicho']
    ) * 100:.6f} %\n\n"
    f"AD p = {fmt_p(
        sup_ols['p_AD'] if grado_final == 1
        else resultado_polinomio['supuestos']['p_AD']
    )}\n"
    f"DW = {(
        sup_ols['DW'] if grado_final == 1
        else resultado_polinomio['supuestos']['DW']
    ):.4f}\n"
    f"BP p = {fmt_p(
        sup_ols['BP_p_F'] if grado_final == 1
        else resultado_polinomio['supuestos']['BP_p_F']
    )}"
)

info.text(
    0.02, 0.86, texto_info,
    fontsize=10.5, va="top", linespacing=1.25
)

fig.suptitle(
    "LINEALIDAD — CURVA DE CALIBRACIÓN",
    fontsize=20, fontweight="bold"
)

plt.show()


# ======================================================================
# 19. GRÁFICA 2 — PANEL DE DIAGNÓSTICO, COMO SCRIPT INICIAL
# ======================================================================

diag_final_completo = (
    diag_ols
    if grado_final == 1
    else resultado_polinomio["diagnosticos"]
)

# La función de diagnóstico devuelve una tupla:
# (tabla_diagnostico, umbral_resid, umbral_leverage, umbral_cook).
# Para las gráficas se utiliza únicamente la tabla.
diag_final = diag_final_completo[0]
sup_final = (
    sup_ols
    if grado_final == 1
    else resultado_polinomio["supuestos"]
)

fig, axes = plt.subplots(2, 2, figsize=(14, 9))

# Q-Q Plot
ax = axes[0, 0]
(theoretical, ordered), (slope_qq, intercept_qq, _) = stats.probplot(
    diag_final["Resid"], dist="norm"
)
ax.scatter(
    theoretical, ordered, s=45,
    color="#1F77B4",
    edgecolor="white", linewidth=0.7
)
line_x = np.array([theoretical.min(), theoretical.max()])
ax.plot(
    line_x,
    intercept_qq + slope_qq * line_x,
    color="#E74C3C",
    linewidth=2
)
ax.set_title(
    f"Q-Q Plot — Normalidad\nAD p = {fmt_p(sup_final['p_AD'])}",
    fontweight="bold"
)
ax.set_xlabel("Cuantiles teóricos")
ax.set_ylabel("Residuales ordenados")
ax.grid(alpha=0.18)

# Residuales vs ajustados
ax = axes[0, 1]
ax.scatter(
    diag_final["Ajuste"],
    diag_final["Resid"],
    s=48, color="#E74C3C",
    edgecolor="white", linewidth=0.7
)
ax.axhline(0, color="#34495E", linestyle="--", linewidth=1.7)
ax.set_title(
    f"Residuales vs ajustados\n"
    f"Breusch-Pagan p = {fmt_p(sup_final['BP_p_F'])}",
    fontweight="bold"
)
ax.set_xlabel("Valores ajustados")
ax.set_ylabel("Residual")
ax.grid(alpha=0.18)

# Residuales vs X
ax = axes[1, 0]
ax.scatter(
    diag_final["X"],
    diag_final["Resid"],
    s=48, color="#2CA02C",
    edgecolor="white", linewidth=0.7
)
ax.axhline(0, color="#34495E", linestyle="--", linewidth=1.7)

media_residual = (
    pd.DataFrame({
        "X": diag_final["X"],
        "Resid": diag_final["Resid"]
    })
    .groupby("X")["Resid"]
    .mean()
    .reset_index()
)

ax.plot(
    media_residual["X"],
    media_residual["Resid"],
    color="#F39C12",
    marker="D",
    linewidth=1.8,
    markersize=6,
    label="Media residual por nivel"
)
ax.set_title("Residuales vs X", fontweight="bold")
ax.set_xlabel(str(col_x))
ax.set_ylabel("Residual")
ax.grid(alpha=0.18)
ax.legend()

# Residuales vs orden
ax = axes[1, 1]
ax.plot(
    diag_final["Obs"],
    diag_final["Resid"],
    color="#8E44AD",
    linewidth=1.5,
    alpha=0.75
)
ax.scatter(
    diag_final["Obs"],
    diag_final["Resid"],
    s=45,
    color="#8E44AD",
    edgecolor="white",
    linewidth=0.7,
    zorder=3
)
ax.axhline(0, color="#34495E", linestyle="--", linewidth=1.7)
ax.set_title(
    f"Residuales vs orden\nDurbin-Watson = {sup_final['DW']:.4f}",
    fontweight="bold"
)
ax.set_xlabel("Orden de observación")
ax.set_ylabel("Residual")
ax.grid(alpha=0.18)

fig.suptitle(
    "Diagnóstico gráfico de los supuestos del modelo",
    fontsize=18,
    fontweight="bold"
)
fig.tight_layout(rect=[0, 0, 1, 0.95])
plt.show()


# ======================================================================
# 20. GRÁFICA 3 — HISTOGRAMA DE RESIDUALES, COMO SCRIPT INICIAL
# ======================================================================

fig, ax = plt.subplots(figsize=(10, 5.8))

res = np.asarray(diag_final["Resid"])
ax.hist(
    res,
    bins="auto",
    density=True,
    color="#16A085",
    edgecolor="white",
    linewidth=1.0,
    alpha=0.82,
    label="Residuales"
)

media_res = np.mean(res)
sd_res = np.std(res, ddof=1)

if sd_res > 0:
    x_normal = np.linspace(res.min(), res.max(), 500)
    ax.plot(
        x_normal,
        stats.norm.pdf(x_normal, media_res, sd_res),
        color="#E67E22",
        linewidth=2.5,
        label="Normal ajustada"
    )

ax.axvline(0, color="#34495E", linestyle="--", linewidth=1.4)
ax.set_title(
    f"Histograma de residuales — "
    f"Anderson-Darling p = {fmt_p(sup_final['p_AD'])}",
    fontsize=16,
    fontweight="bold"
)
ax.set_xlabel("Residual")
ax.set_ylabel("Densidad")
ax.grid(axis="y", alpha=0.18)
ax.spines["top"].set_visible(False)
ax.spines["right"].set_visible(False)
ax.legend()
plt.tight_layout()
plt.show()


# ======================================================================
# 21. GRÁFICA 4 — OBSERVACIONES POCO COMUNES, COMO SCRIPT INICIAL
# ======================================================================

fig, axes = plt.subplots(1, 2, figsize=(14, 5.7))

ax = axes[0]
ax.scatter(
    diag_final["Obs"],
    diag_final["Resid est"],
    s=52,
    color="#3498DB",
    edgecolor="white",
    linewidth=0.8
)
ax.axhline(0, color="#34495E", linestyle="--", linewidth=1.2)
ax.axhline(2, color="#F39C12", linestyle=":", linewidth=1.5)
ax.axhline(-2, color="#F39C12", linestyle=":", linewidth=1.5)
ax.axhline(3, color="#E74C3C", linestyle="--", linewidth=1.2)
ax.axhline(-3, color="#E74C3C", linestyle="--", linewidth=1.2)
ax.set_title("Residuales estandarizados", fontweight="bold")
ax.set_xlabel("Observación")
ax.set_ylabel("Residual estandarizado")
ax.grid(alpha=0.18)

ax = axes[1]
markerline, stemlines, baseline = ax.stem(
    diag_final["Obs"],
    diag_final["Cook"],
    basefmt=" "
)
plt.setp(markerline, color="#E74C3C")
plt.setp(stemlines, color="#E74C3C")
ax.axhline(
    diag_final[0] if False else (
        4 / len(x)
    ),
    color="#8E44AD",
    linestyle="--",
    linewidth=1.5,
    label=f"4/n = {4/len(x):.4f}"
)
ax.set_title("Distancia de Cook", fontweight="bold")
ax.set_xlabel("Observación")
ax.set_ylabel("Cook's Distance")
ax.grid(alpha=0.18)
ax.legend()

fig.suptitle(
    "Diagnóstico de observaciones potencialmente influyentes",
    fontsize=17,
    fontweight="bold"
)
fig.tight_layout(rect=[0, 0, 1, 0.94])
plt.show()



# ======================================================================
# 23. CONCLUSIÓN
# ======================================================================

print("\n" + "=" * 95)
print("12. CONCLUSIÓN")
print("=" * 95)

if lof_ols["evaluable"]:
    if lof_ols["p_LOF"] >= ALPHA:
        print(
            "\n✓ El modelo lineal NO presenta evidencia estadística "
            "de falta de ajuste significativa."
        )
        print("La forma lineal es funcionalmente adecuada según LOF.")
    else:
        print(
            "\n⚠ El modelo lineal presenta falta de ajuste significativa."
        )
        if resultado_polinomio is not None:
            print(
                f"Se seleccionó {nombre_modelo_final} como Plan B."
            )
        else:
            print(
                "Los modelos polinomiales evaluados no cumplieron "
                "los criterios funcionales."
            )
else:
    print(
        "\n⚠ Lack-of-Fit no fue evaluable por falta de réplicas "
        "suficientes."
    )

problemas = []

if sup_final["p_AD"] < ALPHA:
    problemas.append("normalidad")

if sup_final["BP_p_F"] < ALPHA:
    problemas.append("homocedasticidad")

if not (1.5 <= sup_final["DW"] <= 2.5):
    problemas.append("independencia/autocorrelación")

if problemas:
    print(
        "\n⚠ Supuestos que requieren revisión: "
        + ", ".join(problemas) + "."
    )
    print(
        "Estos incumplimientos NO se interpretan automáticamente "
        "como falta de linealidad."
    )
else:
    print(
        "\n✓ No se detectaron señales importantes en los "
        "supuestos evaluados."
    )


# ======================================================================
# 24. TABLA FINAL
# ======================================================================

p_lof_final = (
    lof_ols["p_LOF"]
    if grado_final == 1
    else resultado_polinomio["lof"]["p_LOF"]
)

tabla_final = pd.DataFrame({
    "Criterio": [
        "Pendiente / modelo significativo",
        "Lack-of-Fit",
        "R²",
        "R² ajustado",
        "R² predicho",
        "Normalidad AD",
        "Homocedasticidad Breusch-Pagan",
        "Independencia Durbin-Watson",
        "Modelo final"
    ],
    "Resultado": [
        fmt_p(
            p_pendiente
            if grado_final == 1
            else modelo_final.f_pvalue
        ),
        fmt_p(p_lof_final),
        modelo_final.rsquared,
        modelo_final.rsquared_adj,
        (
            res_ols["r2_predicho"]
            if grado_final == 1
            else resultado_polinomio["r2_predicho"]
        ),
        sup_final["p_AD"],
        sup_final["BP_p_F"],
        sup_final["DW"],
        nombre_modelo_final
    ],
    "Criterio / referencia": [
        "p < 0.05",
        "p ≥ 0.05",
        "Sin umbral universal",
        "Informativo",
        "Informativo",
        "p ≥ 0.05",
        "p ≥ 0.05",
        "1.5 ≤ DW ≤ 2.5",
        "—"
    ]
})

print("\n" + "=" * 95)
print("13. TABLA FINAL DE EVALUACIÓN")
print("=" * 95)
print(tabla_final.to_string(index=False))


# ======================================================================
# 25. EXPORTACIÓN A EXCEL
# ======================================================================

try:
    coef_final = pd.DataFrame({
        "Parámetro": [
            f"β{i}" for i in range(len(modelo_final.params))
        ],
        "Estimación": modelo_final.params,
        "EE": modelo_final.bse,
        "t": modelo_final.tvalues,
        "p": modelo_final.pvalues
    })

    residuales_export = diag_final.copy()

    sensibilidad_export = pd.DataFrame({
        "X": x_grid,
        "Sensibilidad_dY_dX": sens_grid
    })

    tabla_supuestos_final = pd.DataFrame({
        "Prueba": [
            "Anderson-Darling",
            "Durbin-Watson",
            "Breusch-Pagan LM",
            "Breusch-Pagan F"
        ],
        "Estadístico": [
            sup_final["AD"],
            sup_final["DW"],
            sup_final["BP_LM"],
            sup_final["BP_F"]
        ],
        "p-valor": [
            sup_final["p_AD"],
            np.nan,
            sup_final["BP_p_LM"],
            sup_final["BP_p_F"]
        ]
    })

    with pd.ExcelWriter(
        ARCHIVO_SALIDA,
        engine="openpyxl"
    ) as writer:

        base.to_excel(writer, sheet_name="Datos", index=False)
        tabla_coef_ols.to_excel(
            writer, sheet_name="Coef_OLS", index=False
        )
        tabla_anova_ols.to_excel(
            writer, sheet_name="ANOVA_OLS", index=False
        )

        if tabla_lof_ols is not None:
            tabla_lof_ols.to_excel(
                writer, sheet_name="LOF_OLS", index=False
            )

        tabla_criterios_ols.to_excel(
            writer, sheet_name="Criterios_OLS", index=False
        )

        residuales_export.to_excel(
            writer, sheet_name="Diagnostico_Final", index=False
        )

        tabla_supuestos_final.to_excel(
            writer, sheet_name="Supuestos_Final", index=False
        )

        tabla_final.to_excel(
            writer, sheet_name="Evaluacion_Final", index=False
        )

        sensibilidad_export.to_excel(
            writer, sheet_name="Sensibilidad", index=False
        )

        if resultado_polinomio is not None:
            grado = resultado_polinomio["grado"]
            modelo_poly = resultado_polinomio["modelo"]

            pd.DataFrame({
                "Parámetro": [
                    f"β{i}" for i in range(len(modelo_poly.params))
                ],
                "Estimación": modelo_poly.params,
                "EE": modelo_poly.bse,
                "t": modelo_poly.tvalues,
                "p": modelo_poly.pvalues
            }).to_excel(
                writer,
                sheet_name=f"Polinomio_G{grado}",
                index=False
            )

    # Formato básico del Excel
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = load_workbook(ARCHIVO_SALIDA)

        for ws in wb.worksheets:
            ws.freeze_panes = "A2"

            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    "solid", fgColor="1565C0"
                )
                cell.alignment = Alignment(horizontal="center")

            for columna in ws.columns:
                max_len = 0
                letra = columna[0].column_letter

                for celda in columna:
                    if celda.value is not None:
                        max_len = max(
                            max_len, len(str(celda.value))
                        )

                ws.column_dimensions[letra].width = min(
                    max_len + 2, 28
                )

        wb.save(ARCHIVO_SALIDA)

    except Exception as e:
        print(
            f"\nAdvertencia: no se pudo aplicar formato avanzado "
            f"al Excel: {e}"
        )

    print("\n" + "=" * 95)
    print("ANÁLISIS COMPLETADO")
    print("=" * 95)
    print(f"\nResultados exportados a:\n{ARCHIVO_SALIDA}")

except Exception as e:
    print(
        "\nNo se pudo exportar el Excel de resultados."
        f"\nDetalle: {e}"
    )
